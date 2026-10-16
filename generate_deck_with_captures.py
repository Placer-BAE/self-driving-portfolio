"""
중간발표용 PPT 생성 (웹페이지 실제 캡쳐본 4종 삽입 및 선배 피드백/기초데이터/21개 모델 전면 반영)
- 16:9 와이드스크린 (13.333" x 7.5")
- 맥킨지/BCG 스타일 엄격한 개조식 명사형 종결
- 선배 피드백 3대 개선안 (AI Co-Pilot 정의, 타깃 변동성 매칭, 기관 실무 제약조건)
- 야후 파이낸스 159개월 실데이터 & Exhibit 5 기반 21개 후보군 명확화
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = r"C:\Users\bgh68\.gemini\antigravity\scratch\self-driving-portfolio\The_Self_Driving_Portfolio_중간발표_최종본.pptx"
OUT_LEGACY = r"C:\Users\bgh68\.gemini\antigravity\scratch\self-driving-portfolio\The_Self_Driving_Portfolio_중간발표_간결본.pptx"

IMG_DIR = r"C:\Users\bgh68\.gemini\antigravity\brain\362c0791-8de4-4bb2-af50-e049d89fa8db\.user_uploaded"
IMG1 = os.path.join(IMG_DIR, "media_1791889036964.png") # 메인 대시보드 + 스테퍼 + KPI
IMG2 = os.path.join(IMG_DIR, "media_1791889094889.png") # 대자산군 도넛 + 60/40 비교
IMG3 = os.path.join(IMG_DIR, "media_1791889106342.png") # 18개 자산군 비중 표
IMG4 = os.path.join(IMG_DIR, "media_1791889122698.png") # 전략 심의 보르다 투표

NAVY = RGBColor(5, 28, 44)
BLUE = RGBColor(0, 82, 204)
BLUE_BG = RGBColor(239, 246, 255)
GRAY_BG = RGBColor(246, 248, 250)
LINE = RGBColor(214, 220, 228)
DARK = RGBColor(15, 23, 42)
BODY = RGBColor(71, 85, 105)
MUTED = RGBColor(148, 163, 184)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(5, 150, 105)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(185, 28, 28)
FONT = "Malgun Gothic"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh

def text(s, x, y, w, h, t, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = ln
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb

def header(s, section, title, page):
    box(s, 0, 0, 13.333, 0.08, fill=NAVY)
    text(s, 0.6, 0.35, 10, 0.3, section, 11, True, BLUE)
    text(s, 0.6, 0.65, 12, 0.6, title, 24, True, NAVY)
    box(s, 0.6, 1.35, 12.13, 0.015, fill=LINE)
    text(s, 12.2, 7.0, 0.6, 0.3, str(page), 10, color=MUTED, align=PP_ALIGN.RIGHT)

def card(s, x, y, w, h, title, items, accent=BLUE, fill=GRAY_BG, title_size=16, item_size=13, spacing=7):
    box(s, x, y, w, h, fill=fill)
    box(s, x, y, w, 0.06, fill=accent)
    text(s, x + 0.22, y + 0.22, w - 0.44, 0.45, title, title_size, True, NAVY)
    text(s, x + 0.22, y + 0.78, w - 0.44, h - 0.88, ["·  " + i for i in items], item_size, color=BODY, spacing=spacing)

def add_image_card(s, x, y, w, h, img_path):
    box(s, x - 0.04, y - 0.04, w + 0.08, h + 0.08, fill=WHITE, line=LINE, lw=1)
    s.shapes.add_picture(img_path, Inches(x), Inches(y), Inches(w), Inches(h))

# ================================================================ 1. 표지
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
box(s, 0, 0, 0.25, 7.5, fill=BLUE)
text(s, 1.0, 2.0, 11, 0.5, "중간발표", 16, True, RGBColor(147, 197, 253))
text(s, 1.0, 2.5, 11, 1.0, "The Self-Driving Portfolio", 44, True, WHITE)
text(s, 1.0, 3.55, 11, 0.6, "투자위원회 의사결정 지원을 위한 AI Co-Pilot (Decision Support System)", 22, False, RGBColor(226, 232, 240))
box(s, 1.0, 4.45, 1.2, 0.03, fill=BLUE)
text(s, 1.0, 4.75, 11, 0.4, "Base Paper  |  Andrew Ang, Nazym Azimbayev, Andrey Kim (2026)", 13, color=MUTED)
text(s, 1.0, 5.15, 11, 0.4, "Live Demo   |  self-driving-portfolio.onrender.com", 13, color=RGBColor(147, 197, 253))

# ================================================================ 2. 목차
s = prs.slides.add_slide(BLANK)
header(s, "AGENDA", "목차", 2)
agenda = [("01", "논문 배경 및 포지셔닝", "기관 SAA 구조적 한계 · 6단계 파이프라인 · AI Co-Pilot 정의"),
          ("02", "주제 선정 이유", "시의성 · 실무 거버넌스 부합성 · 구현 및 검증 가능성"),
          ("03", "웹 시스템 구현 및 실화면", "159개월 실데이터 · 21개 모델 후보군 · 4대 실화면 심층 분석"),
          ("04", "중간 성과 및 향후 로드맵", "타깃 변동성 매칭(Vol Scaling) · 실무 제약조건 · 기말 계획")]
for i, (n, t, d) in enumerate(agenda):
    y = 1.9 + i * 1.25
    text(s, 1.2, y, 1.2, 0.8, n, 36, True, BLUE)
    text(s, 2.6, y + 0.05, 5.5, 0.5, t, 22, True, NAVY)
    text(s, 2.6, y + 0.6, 8.5, 0.4, d, 14, color=BODY)
    if i < 3:
        box(s, 2.6, y + 1.1, 9.5, 0.01, fill=LINE)

# ================================================================ 3. 논문 배경: 한계
s = prs.slides.add_slide(BLANK)
header(s, "01  논문 배경", "기관 SAA의 구조적 한계 및 Co-Pilot 필요성", 3)
problems = [("연 1회", "경직된 주기", ["시장 국면 변화 대응 지연", "형식적 재검토"]),
            ("사일로", "부서 간 단절", ["부서별 전제 불일치", "수작업 취합 비용"]),
            ("사후 검증", "리스크 통제", ["확정 후 리스크 검토", "수정 비용 과다"]),
            ("기록 부재", "감사 추적", ["조정 사유 불투명", "재현 · 검증 곤란"])]
for i, (big, sub, items) in enumerate(problems):
    x = 0.6 + i * 3.08
    box(s, x, 1.8, 2.88, 3.4, fill=GRAY_BG)
    text(s, x + 0.25, 2.05, 2.4, 0.7, big, 28, True, RED)
    text(s, x + 0.25, 2.8, 2.4, 0.4, sub, 15, True, NAVY)
    text(s, x + 0.25, 3.4, 2.45, 1.7, ["·  " + it for it in items], 13, color=BODY, spacing=8)
box(s, 0.6, 5.55, 12.13, 1.05, fill=NAVY)
text(s, 0.95, 5.55, 11.5, 1.05, "핵심 정의  →  인간 CIO 대체가 아닌, 투자위원회의 시나리오 분석 및 SAA 의사결정을 보좌하는 AI Co-Pilot",
     17, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 4. 6단계 파이프라인
s = prs.slides.add_slide(BLANK)
header(s, "01  논문 배경", "6단계 에이전트 파이프라인 아키텍처", 4)
stages = [("거시 국면", "Macro (1)", "성장·물가 4분면 진단"),
          ("자산군 CMA", "CMA (18)", "18개 자산 10년 예측"),
          ("공분산 행렬", "Covariance (1)", "Ledoit-Wolf / OAS 수축"),
          ("포트폴리오 구성", "PC (21)", "4개 계열 21개 모델 경쟁"),
          ("전략 심의 · 투표", "CRO + 21 심사역", "수정 보르다 카운트 투표"),
          ("최종 앙상블", "CIO (1)", "역추적오차(TE) 결합")]
w = 2.0
for i, (t, en, d) in enumerate(stages):
    x = 0.6 + i * 2.03
    fill = BLUE if i == 4 else NAVY
    box(s, x, 2.0, w + 0.15, 1.3, fill=fill, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    text(s, x + 0.35, 2.0, w - 0.45, 1.3, str(i + 1), 26, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x, 3.6, w, 0.4, t, 15, True, NAVY, PP_ALIGN.CENTER)
    text(s, x, 4.05, w, 0.35, en, 11.5, True, BLUE, PP_ALIGN.CENTER)
    text(s, x, 4.45, w, 0.4, d, 11.5, color=BODY, align=PP_ALIGN.CENTER)
box(s, 6.7, 5.25, 4.1, 0.55, fill=BLUE_BG, line=BLUE)
text(s, 6.7, 5.25, 4.1, 0.55, "↺  CRO IPS 위반 시 4단계 재최적화", 13, True, BLUE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.3, 12, 0.4, "입력: 18개 ETF 실데이터 159개월 · 거시지표   |   출력: 최종 권고 비중 · 투자위원회 안건 메모", 13, color=BODY)

# ================================================================ 5. 4대 설계 원칙
s = prs.slides.add_slide(BLANK)
header(s, "01  논문 배경", "기관 거버넌스를 위한 4대 설계 원칙", 5)
principles = [("전문화", "Specialization", "거시 · 자산군 · 최적화 · 리스크 에이전트 분업"),
              ("도구 결정론", "Tool Determinism", "LLM = 정성 해석 · 위원회 토론\nPython = 수학 행렬 계산 (환각 0%)"),
              ("적대적 검증", "Adversarial Critique", "CRO 거부권(Veto) → IPS 미달 배제"),
              ("감사 추적", "Audit Trail", "보르다 투표 · 소수의견 · 심의 사유 기록")]
for i, (k, en, d) in enumerate(principles):
    x = 0.6 + i * 3.08
    box(s, x, 1.9, 2.88, 3.9, fill=BLUE_BG if i == 1 else GRAY_BG)
    text(s, x + 0.25, 2.2, 0.8, 0.6, f"0{i+1}", 18, True, BLUE)
    text(s, x + 0.25, 2.85, 2.4, 0.6, k, 24, True, NAVY)
    text(s, x + 0.25, 3.5, 2.4, 0.4, en, 12, True, BLUE)
    box(s, x + 0.25, 4.0, 0.6, 0.02, fill=BLUE)
    text(s, x + 0.25, 4.2, 2.4, 1.4, d.split("\n"), 13.5, color=BODY, spacing=6)
text(s, 0.6, 6.3, 12, 0.4, "핵심  →  LLM 수치 환각 원천 차단 (Tool Determinism) + 투명한 이사회 감사 추적성", 15, True, NAVY)

# ================================================================ 6. 주제 선정 이유
s = prs.slides.add_slide(BLANK)
header(s, "02  주제 선정 이유", "왜 이 논문인가: 학술 연구에서 실무 시스템으로", 6)
reasons = [("시의성", ["2026년 4월 최신 연구", "Agentic AI × 자산배분 접목", "기관 AI Co-Pilot 도입 선도"]),
           ("실무성", ["연기금 실제 SAA/DAA 의사결정", "투자위원회 · IPS 거버넌스 완비", "수리 최적화와 정성 토론 결합"]),
           ("구현 가능성", ["수학 공식 완비 (OAS · BL · CVaR)", "18개 ETF 실데이터 159개월 검증", "24/7 웹 시연으로 실효성 입증"])]
for i, (t, items) in enumerate(reasons):
    x = 0.6 + i * 4.1
    card(s, x, 1.8, 3.9, 3.0, t, items, title_size=22, item_size=15)
box(s, 0.6, 5.3, 12.13, 0.9, fill=NAVY)
text(s, 0.95, 5.3, 11.5, 0.9, "목표  →  논문의 50개 에이전트 아키텍처를 실제 데이터 기반의 실동작 웹 플랫폼으로 재현",
     18, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 7. 시스템 구성도 및 기초데이터
s = prs.slides.add_slide(BLANK)
header(s, "03  웹 시스템 구현", "시스템 아키텍처 및 기초데이터 체계", 7)
blocks = [("데이터 계층 (실데이터)", [
              "야후 파이낸스 실데이터 159개월",
              "(2013-07 ~ 2026-09 수정 총수익률)",
              "18개 대표 ETF 유동성 대용치",
              "FRED 실시간 거시지표 (성장·물가·금리)"
          ]),
          ("수리 엔진 (백엔드)", [
              "Python · FastAPI 프레임워크",
              "NumPy · SciPy 결정론적 행렬 연산",
              "Exhibit 5 4대 계열 21개 모델 최적화",
              "OAS 공분산 수축 및 보르다 투표 집계"
          ]),
          ("대시보드 (프론트엔드)", [
              "React 19 · TypeScript · Tailwind CSS",
              "거버넌스 L3 상태 및 6단계 스테퍼",
              "자산배분 탭 · 보르다 투표 · 진단 차트",
              "IPS 파라미터 즉각 재실행 인터페이스"
          ]),
          ("배포 및 인프라", [
              "Docker 경량 멀티 스테이지 빌드",
              "Render 클라우드 24/7 무중단 호스팅",
              "GitHub Actions CI/CD 완전 자동화",
              "반응형 웹 UI (데스크톱/모바일)"
          ])]
for i, (t, items) in enumerate(blocks):
    x = 0.6 + i * 3.14
    card(s, x, 1.9, 2.7, 3.3, t, items, accent=BLUE if i != 1 else NAVY, title_size=17, item_size=12, spacing=5)
    if i < 3:
        text(s, x + 2.72, 3.2, 0.4, 0.6, "→", 26, True, BLUE, PP_ALIGN.CENTER)
box(s, 0.6, 5.5, 12.13, 0.9, fill=BLUE_BG, line=BLUE)
text(s, 0.8, 5.55, 11.7, 0.8, [
    "·  데이터 투명성: 상용 벤더가 아닌 야후 파이낸스 159개월 실데이터(18개 ETF)로 완전 공개 검증 가능",
    "·  Tool Determinism: 수치 연산은 100% 파이썬 코드가 전담하여 환각 0% 보장, LLM은 거버넌스 해설 전담"
], 12.5, color=DARK, spacing=3)

# ================================================================ 8. [캡쳐 1] 메인 대시보드
s = prs.slides.add_slide(BLANK)
header(s, "03  웹 시스템 구현", "메인 대시보드 및 핵심 거버넌스 지표 (실화면)", 8)
card(s, 0.6, 1.6, 4.2, 5.0, "거버넌스 L3 & 핵심 성과 지표", [
    "거버넌스 L3: 투자위원회 지원 AI Co-Pilot",
    "IPS 적격 심사: 21개 후보 중 2개만 최종 통과",
    "예상 명목수익률: 5.75% (실질수익률 +3.3%)",
    "사전 연간변동성: 8.85% (IPS [8~12%] 충족)",
    "샤프 지수: 0.231 (vs 60/40 0.219 대비 우수)",
    "예상 최대낙폭: -18.59% (60/40 -34.3% 대비 방어)",
    "사전 트래킹 에러: 0.77% (IPS < 6.0% 준수)",
    "유효 자산수: 10.2개 (충분한 자산 다변화)"
], title_size=18, item_size=12.5)
add_image_card(s, 5.1, 1.6, 7.6, 5.0, IMG1)

# ================================================================ 9. [캡쳐 2] 대자산군 배분 및 벤치마크
s = prs.slides.add_slide(BLANK)
header(s, "03  웹 시스템 구현", "대자산군 배분 및 60/40 벤치마크 비교 (실화면)", 9)
card(s, 0.6, 1.6, 4.2, 5.0, "대자산군 배분 및 벤치마크 진단", [
    "주식군 (Equity): 56.49% (대형·성장·가치·중소)",
    "채권군 (Fixed Income): 40.71% (중기·단기·장기·회사채)",
    "실물자산 (Real Assets): 1.65% (금·원자재·리츠)",
    "현금 버퍼 (Cash): 1.17% (초단기국채 유동성)",
    "거시 국면 대응: 스태그플레이션 방어 주식 언더웨이트",
    "채권 듀레이션 분산 및 현금 확보로 완충력 강화",
    "최대 낙폭(MDD): -34.3% → -18.59%로 대폭 개선",
    "사전 트래킹 에러: 0.77% (벤치마크 추종 균형)"
], title_size=18, item_size=12.5)
add_image_card(s, 5.1, 1.6, 7.6, 5.0, IMG2)

# ================================================================ 10. [캡쳐 3] 18개 자산군별 비중
s = prs.slides.add_slide(BLANK)
header(s, "03  웹 시스템 구현", "18개 자산군별 최종 권고 비중 및 21개 후보군 (Exhibit 11)", 10)
card(s, 0.6, 1.6, 4.2, 5.0, "후보군 21개 모델 및 최종 비중", [
    "후보군: 논문 Exhibit 5 기반 4대 계열 21개 모델",
    "· A: 휴리스틱 5 (60/40, 동일가중, 모멘텀 등)",
    "· B: 수익률최적화 5 (MVO, 블랙-리터만 등)",
    "· C: 위험구조화 5 (위험균형, 최소분산, 최대분산 등)",
    "· D: 비전통 6 (TPA, 적대적분산, 최대엔트로피 등)",
    "미국 대형주 (SPY): 28.56% (위험 기여 45.2%)",
    "미국 중기국채 (IEF): 16.57% (포트폴리오 핵심 안전판)",
    "스타일 주식 (VUG/VTV): 각 9.41% 균형 배분",
    "단기/장기채 (SHY/TLT): 각 5.09% | 회사채: 5.09%"
], title_size=17, item_size=12)
add_image_card(s, 5.1, 1.6, 7.6, 5.0, IMG3)

# ================================================================ 11. [캡쳐 4] 전략 심의 보르다 투표
s = prs.slides.add_slide(BLANK)
header(s, "03  웹 시스템 구현", "전략 심의: 수정 보르다 투표 및 IPS 적격 필터링 (Exhibit 9)", 11)
card(s, 0.6, 1.6, 4.2, 5.0, "전략 심의 및 적격 필터링 진단", [
    "평가 점수: 동료 투표 55% + CRO 정량 점수 45%",
    "상위 모델의 IPS 변동성 하한(8%) 미달 탈락:",
    "· 1위 최대 분산 (2.7%), 2위 BL (6.9%), 3위 위험균형 (6.5%)",
    "4위 TPA (변동성 8.3%), 6위 시총가중 (9.2%)만 통과",
    "엄격한 IPS 거버넌스 시연 vs 우수 알파 모델 누락",
    "기말 해결책: 타깃 변동성 매칭(Vol Scaling) 도입",
    "적대적 분산 에이전트: 소수파 관점 및 편향 견제 기록",
    "최종 앙상블: 적격 2개 모델 역추적오차(TE) 결합"
], title_size=17, item_size=12)
add_image_card(s, 5.1, 1.6, 7.6, 5.0, IMG4)

# ================================================================ 12. 현재 이슈 및 선배 피드백 반영 과제
s = prs.slides.add_slide(BLANK)
header(s, "04  향후 계획", "현재 이슈 및 선배 피드백 반영 검토 과제", 12)
issues = [("타깃 변동성 매칭\n(Vol Scaling)", [
              "문제점: BL·위험균형 등 우수 모델의 변동성 하한(8%) 미달 탈락",
              "개선안: 현금/레버리지 조절로 목표 변동성(8~12%) 스케일링",
              "기대효과: 우수 알파 모델 복원 및 앙상블 다변화 극대화"
          ]),
          ("실무 제약조건 추가\n(Turnover & Cost)", [
              "문제점: 정적 자산배분 가정으로 리밸런싱 마찰비용 미반영",
              "개선안: 연간 회전율(Turnover) 상한 및 슬리피지·수수료 모델링",
              "기대효과: 거래비용 차감 순수익률(Net Return) 기반 실무 검증"
          ]),
          ("원/달러 환헤지\n(FX Hedge) 모듈", [
              "문제점: 해외 ETF(미 달러) 100% 노출에 따른 환율 위험 미통제",
              "개선안: 국내 연기금 가이드라인 기반 환헤지 비율(0~100%) 반영",
              "기대효과: 환변동성 리스크 버퍼 구축 및 국내 기관 정합성 확보"
          ]),
          ("AI Co-Pilot 거버넌스\n포지셔닝 확립", [
              "문제점: 완전 자율 운용으로 오인 시 연기금 수탁자 책임 저촉",
              "개선안: 투자위원회 시나리오 분석 및 SAA 의사결정 보좌 DSS",
              "기대효과: 인간 위원회와 협업하는 현실적 거버넌스 모델 확립"
          ])]
for i, (t, items) in enumerate(issues):
    x = 0.6 + i * 3.08
    card(s, x, 1.8, 2.88, 4.4, t, items, accent=AMBER, title_size=16, item_size=12, spacing=6)

# ================================================================ 13. 기말 로드맵
s = prs.slides.add_slide(BLANK)
header(s, "04  향후 계획", "선배 피드백을 반영한 기말 로드맵", 13)
phases = [("Phase 1", "1~2주", "변동성 매칭 & 모델 고도화", [
              "타깃 변동성 매칭(Vol Scaling) 구현",
              "블랙-리터만 · 위험균형 모델 적격 통과",
              "21개 최적화 모델 파이프라인 수치 안정화"
          ]),
          ("Phase 2", "3~4주", "기관 실무 제약조건 반영", [
              "원/달러 환헤지(FX Hedge) 룰 엔진 개발",
              "회전율(Turnover) 제약식 및 거래비용 반영",
              "FRED 거시지표 API 실시간 동기화"
          ]),
          ("Phase 3", "5~6주", "LLM Co-Pilot 실연동", [
              "거시 · CIO 실 LLM API(Gemini/Claude) 연동",
              "투자위원회 안건 및 시나리오 메모 자동 생성",
              "인간 위원회-AI Co-Pilot 협업 워크플로우"
          ]),
          ("Phase 4", "7~8주", "실증 백테스트 & 최종 보고", [
              "거래비용 반영 순수익률 다구간 백테스트",
              "60/40 및 단일 모델 대비 샤프·MDD 심층 비교",
              "기말 종합 보고서 및 최종 시연 발표"
          ])]
box(s, 0.9, 2.35, 11.6, 0.04, fill=LINE)
for i, (ph, wk, t, items) in enumerate(phases):
    x = 0.6 + i * 3.08
    text(s, x + 0.25, 1.65, 2.6, 0.3, ph, 12, True, BLUE)
    text(s, x + 1.2, 1.6, 1.6, 0.4, wk, 16, True, NAVY)
    box(s, x + 0.25, 2.17, 0.4, 0.4, fill=BLUE if i == 0 else NAVY, shape=MSO_SHAPE.OVAL)
    box(s, x, 2.9, 2.88, 3.2, fill=BLUE_BG if i == 0 else GRAY_BG)
    text(s, x + 0.25, 3.12, 2.4, 0.55, t, 17, True, NAVY)
    text(s, x + 0.25, 3.8, 2.45, 2.2, ["·  " + it for it in items], 13, color=BODY, spacing=10)
text(s, 0.6, 6.6, 12, 0.4, "※ 선배 피드백 3대 권고사항(Vol Scaling, 기관 제약조건, Co-Pilot 확립)을 일정 전반에 전면 반영", 11.5, color=MUTED)

# ================================================================ 14. 기말 산출물 + Q&A
s = prs.slides.add_slide(BLANK)
header(s, "04  향후 계획", "기말 목표 산출물 및 결언", 14)
outs = [("투자위원회 지원 AI Co-Pilot 웹 시스템", "인간 CIO 의사결정을 보좌하는 L3 수준 자율 SAA/DAA 플랫폼 완성"),
        ("환헤지 · 거래비용 반영 기관급 수리 엔진", "21개 모델 수리 최적화 + Vol Scaling + 회전율/비용 제약 완비"),
        ("24/7 글로벌 라이브 데모 & API", "Render 클라우드 상시 운영 및 외부 연동 API 문서화 (onrender.com)"),
        ("국면별 실증 분석 보고서 및 오픈소스", "장기 백테스트 실증 성과 입증 및 GitHub 완전 오픈소스 가이드 제공")]
for i, (k, v) in enumerate(outs):
    x = 0.6 + (i % 2) * 6.15
    y = 1.8 + (i // 2) * 1.75
    box(s, x, y, 5.95, 1.5, fill=GRAY_BG)
    box(s, x, y, 0.08, 1.5, fill=BLUE)
    text(s, x + 0.4, y + 0.28, 5.3, 0.5, k, 19, True, NAVY)
    text(s, x + 0.4, y + 0.85, 5.3, 0.45, v, 13.5, color=BODY)
box(s, 0.6, 5.6, 12.13, 1.1, fill=NAVY)
text(s, 0.6, 5.6, 12.13, 1.1, "Q & A", 30, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("Saved PPT to:", OUT)

try:
    prs.save(OUT_LEGACY)
    print("Saved PPT to:", OUT_LEGACY)
except Exception as e:
    print("Note: Legacy file locked, saved to primary:", OUT)
