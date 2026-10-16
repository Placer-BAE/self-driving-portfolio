"""
발표_대본_최종.md 파일을 가독성 높은 Word (.docx) 문서로 변환하는 스크립트
출력 파일: The_Self_Driving_Portfolio_15분_발표대본_및_QA_최종본.docx
"""

import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "발표_대본_최종.md"
OUT_DOCX = ROOT / "The_Self_Driving_Portfolio_15분_발표대본_및_QA_최종본.docx"

NAVY = RGBColor(5, 28, 44)
BLUE = RGBColor(0, 82, 204)
DARK = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)
RED_ACCENT = RGBColor(185, 28, 28)


def set_run_font(run, font_name="맑은 고딕", size=10.5, bold=False, color=DARK):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_formatted_paragraph(doc, text, size=10.5, space_before=2, space_after=6, left_indent=0, base_color=DARK, is_quote=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.38
    if left_indent > 0:
        p.paragraph_format.left_indent = Inches(left_indent)

    # **bold** 및 `code` 파싱
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            inner = tok[2:-2]
            r = p.add_run(inner)
            set_run_font(r, size=size, bold=True, color=BLUE if not is_quote else NAVY)
        elif tok.startswith("`") and tok.endswith("`"):
            inner = tok[1:-1]
            r = p.add_run(inner)
            set_run_font(r, font_name="Consolas", size=size - 0.5, bold=True, color=RED_ACCENT)
        else:
            r = p.add_run(tok)
            set_run_font(r, size=size, bold=False, color=base_color)
    return p


def main():
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # 여백 설정 (인쇄 및 태블릿 열람 최적화)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    for raw_line in lines:
        line = raw_line.strip()
        if not line or line == "---":
            continue

        # H1 제목 (# )
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(line[2:].strip())
            set_run_font(r, size=18, bold=True, color=NAVY)
            continue

        # H2 파트 제목 (## )
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("■ " + line[3:].strip())
            set_run_font(r, size=14.5, bold=True, color=BLUE)
            continue

        # H3 슬라이드/영역 제목 (### )
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line[4:].strip())
            set_run_font(r, size=12.5, bold=True, color=NAVY)
            continue

        # H4 질문 제목 (#### )
        if line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("▶ " + line[5:].strip())
            set_run_font(r, size=11.5, bold=True, color=RED_ACCENT)
            continue

        # 인용구 (> )
        if line.startswith(">"):
            q_text = line.lstrip(">").strip()
            add_formatted_paragraph(
                doc, q_text, size=10.0, space_before=2, space_after=4,
                left_indent=0.25, base_color=MUTED, is_quote=True
            )
            continue

        # 불릿 (* 또는 - 또는 숫자.)
        if line.startswith("* ") or line.startswith("- "):
            b_text = "• " + line[2:].strip()
            add_formatted_paragraph(
                doc, b_text, size=10.5, space_before=2, space_after=4,
                left_indent=0.18, base_color=DARK
            )
            continue

        if re.match(r"^\d+\.\s", line):
            add_formatted_paragraph(
                doc, line, size=10.5, space_before=2, space_after=4,
                left_indent=0.18, base_color=DARK
            )
            continue

        # 일반 본문 대본
        add_formatted_paragraph(doc, line, size=10.5, space_before=3, space_after=7, base_color=DARK)

    doc.save(OUT_DOCX)
    print("Saved Word file to:", OUT_DOCX)


if __name__ == "__main__":
    main()
