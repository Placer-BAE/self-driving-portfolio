"""
중간발표용 간결 PPT 생성 (단어 위주, 줄글 최소화)
구성: 논문 배경 → 주제 선정 이유 → 구현 과정/현황 → 이슈 → 기말까지 계획
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = r"C:\Users\bgh68\.gemini\antigravity\scratch\self-driving-portfolio\The_Self_Driving_Portfolio_중간발표_간결본.pptx"

NAVY = RGBColor(5, 28, 44)
BLUE = RGBColor(0, 82, 204)
BLUE_BG = RGBColor(239, 246, 255)
GRAY_BG = RGBColor(246, 248, 250)
LINE = RGBColor(214, 220, 228)
DARK = RGBColor(15, 23, 42)
BODY = RGBColor(71, 85, 105)
MUTED = RGBColor(148, 163, 184)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(5, 150, 105)
AMBER = RGBColor(217, 119, 6)
GREY = RGBColor(148, 163, 184)
RED = RGBColor(185, 28, 28)
FONT = "Malgun Gothic"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, t, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = ln
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def header(s, section, title, page):
    box(s, 0, 0, 13.333, 0.08, fill=NAVY)
    text(s, 0.6, 0.35, 10, 0.3, section, 11, True, BLUE)
    text(s, 0.6, 0.65, 12, 0.6, title, 24, True, NAVY)
    box(s, 0.6, 1.35, 12.13, 0.015, fill=LINE)
    text(s, 12.2, 7.0, 0.6, 0.3, str(page), 10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(s, x, y, w, h, title, items, accent=BLUE, fill=GRAY_BG, title_size=16, item_size=13):
    box(s, x, y, w, h, fill=fill)
    box(s, x, y, w, 0.06, fill=accent)
    text(s, x + 0.25, y + 0.3, w - 0.5, 0.5, title, title_size, True, NAVY)
    text(s, x + 0.25, y + 0.95, w - 0.5, h - 1.1, ["·  " + i for i in items], item_size, color=BODY, spacing=8)


# ---------------------------------------------------------------- 1. 표지
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
box(s, 0, 0, 0.25, 7.5, fill=BLUE)
text(s, 1.0, 2.2, 11, 0.5, "중간발표", 16, True, RGBColor(147, 197, 253))
text(s, 1.0, 2.7, 11, 1.0, "The Self-Driving Portfolio", 44, True, WHITE)
text(s, 1.0, 3.7, 11, 0.6, "멀티 에이전트 기반 기관 전략적 자산배분(SAA) 구현", 22, False, RGBColor(226, 232, 240))
box(s, 1.0, 4.6, 1.2, 0.03, fill=BLUE)
text(s, 1.0, 4.85, 11, 0.4, "Base Paper  |  Ang, Azimbayev, Kim (2026)", 13, color=MUTED)
text(s, 1.0, 5.25, 11, 0.4, "github.com/Placer-BAE/self-driving-portfolio", 13, color=MUTED)

# ---------------------------------------------------------------- 2. 목차
s = prs.slides.add_slide(BLANK)
header(s, "AGENDA", "목차", 2)
agenda = [("01", "논문 배경", "문제의식 · 핵심 구조 · 설계 원칙"),
          ("02", "주제 선정 이유", "시의성 · 실무성 · 구현 가능성"),
          ("03", "구현 과정 및 현황", "시스템 구성 · 진행 상태 · 중간 결과"),
          ("04", "기말까지 계획", "이슈 · 로드맵 · 최종 산출물")]
for i, (n, t, d) in enumerate(agenda):
    y = 1.9 + i * 1.25
    text(s, 1.2, y, 1.2, 0.8, n, 36, True, BLUE)
    text(s, 2.6, y + 0.05, 5, 0.5, t, 22, True, NAVY)
    text(s, 2.6, y + 0.6, 8, 0.4, d, 14, color=BODY)
    if i < 3:
        box(s, 2.6, y + 1.1, 9.5, 0.01, fill=LINE)

# ---------------------------------------------------------------- 3. 논문 배경: 문제의식
s = prs.slides.add_slide(BLANK)
header(s, "01  논문 배경", "기관 SAA의 구조적 한계", 3)
problems = [("연 1회", "경직된 주기", ["시장 국면 변화 대응 지연", "형식적 재검토"]),
            ("사일로", "부서 간 단절", ["부서별 전제 불일치", "수작업 취합"]),
            ("사후 검증", "리스크 통제", ["확정 후 리스크 검토", "수정 비용 과다"]),
            ("기록 부재", "감사 추적", ["조정 사유 불투명", "재현 · 검증 곤란"])]
for i, (big, sub, items) in enumerate(problems):
    x = 0.6 + i * 3.08
    box(s, x, 1.8, 2.88, 3.4, fill=GRAY_BG)
    text(s, x + 0.25, 2.05, 2.4, 0.7, big, 28, True, RED)
    text(s, x + 0.25, 2.8, 2.4, 0.4, sub, 15, True, NAVY)
    text(s, x + 0.25, 3.4, 2.45, 1.7, ["·  " + it for it in items], 13, color=BODY, spacing=8)
box(s, 0.6, 5.55, 12.13, 1.05, fill=NAVY)
text(s, 0.95, 5.55, 11.5, 1.05, "논문 제안  →  AI 에이전트 위원회가 SAA 전 과정을 분업 · 상호 검증 · 기록",
     18, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 4. 논문 핵심 구조
s = prs.slides.add_slide(BLANK)
header(s, "01  논문 배경", "6단계 에이전트 파이프라인", 4)
stages = [("거시 진단", "Macro", "경제 국면 판정"),
          ("수익 · 위험 가정", "CMA · Covariance", "기대수익률 · 공분산"),
          ("포트폴리오 구성", "PC Agents", "복수 최적화 후보"),
          ("리스크 검증", "CRO", "IPS 위반 시 거부"),
          ("위원회 투표", "Committee", "찬반 · 소수의견"),
          ("최종 승인", "CIO", "배분 확정 · 이사회 메모")]
w = 2.0
for i, (t, en, d) in enumerate(stages):
    x = 0.6 + i * 2.03
    fill = BLUE if i == 3 else NAVY
    sh = box(s, x, 2.0, w + 0.15, 1.3, fill=fill, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    text(s, x + 0.35, 2.0, w - 0.45, 1.3, str(i + 1), 26, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x, 3.6, w, 0.4, t, 15, True, NAVY, PP_ALIGN.CENTER)
    text(s, x, 4.05, w, 0.35, en, 12, True, BLUE, PP_ALIGN.CENTER)
    text(s, x, 4.45, w, 0.4, d, 12, color=BODY, align=PP_ALIGN.CENTER)
# 피드백 루프 표시
box(s, 6.7, 5.25, 4.1, 0.55, fill=BLUE_BG, line=BLUE)
text(s, 6.7, 5.25, 4.1, 0.55, "↺  거부 시 3단계 재최적화", 13, True, BLUE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.3, 12, 0.4, "입력: 거시지표 · 자산군 수익률   |   출력: 최종 SAA 비중 · 투자 메모", 13, color=BODY)

# ---------------------------------------------------------------- 5. 설계 원칙
s = prs.slides.add_slide(BLANK)
header(s, "01  논문 배경", "4대 설계 원칙", 5)
principles = [("전문화", "Specialization", "역할별 에이전트 분업"),
              ("도구 결정론", "Tool Determinism", "LLM = 해석 · 토론\n코드 = 수학 계산"),
              ("적대적 검증", "Adversarial Critique", "CRO 거부권 → 재작업"),
              ("감사 추적", "Audit Trail", "투표 · 근거 · 사유 기록")]
for i, (k, en, d) in enumerate(principles):
    x = 0.6 + i * 3.08
    box(s, x, 1.9, 2.88, 3.9, fill=BLUE_BG if i == 1 else GRAY_BG)
    text(s, x + 0.25, 2.2, 0.8, 0.6, f"0{i+1}", 18, True, BLUE)
    text(s, x + 0.25, 2.85, 2.4, 0.6, k, 24, True, NAVY)
    text(s, x + 0.25, 3.5, 2.4, 0.4, en, 12, True, BLUE)
    box(s, x + 0.25, 4.1, 0.6, 0.02, fill=BLUE)
    text(s, x + 0.25, 4.35, 2.4, 1.5, d.split("\n"), 14, color=BODY, spacing=6)
text(s, 0.6, 6.4, 12, 0.4, "핵심  →  LLM 수치 환각 차단 + 의사결정 투명성", 15, True, NAVY)

# ---------------------------------------------------------------- 6. 주제 선정 이유
s = prs.slides.add_slide(BLANK)
header(s, "02  주제 선정 이유", "왜 이 논문인가", 6)
reasons = [("시의성", ["2026년 최신 연구", "Agentic AI × 자산운용", "실무 도입 초기 단계"]),
           ("실무성", ["연기금 · 국부펀드 SAA 프로세스", "위원회 · IPS 등 실제 거버넌스 반영", "금융공학 수업 내용과 직결"]),
           ("구현 가능성", ["명확한 수식 (OAS · BL · CVaR)", "공개 데이터로 재현 가능", "웹 시연으로 결과 확인 용이"])]
for i, (t, items) in enumerate(reasons):
    x = 0.6 + i * 4.1
    card(s, x, 1.8, 3.9, 3.0, t, items, title_size=22, item_size=15)
box(s, 0.6, 5.3, 12.13, 0.9, fill=NAVY)
text(s, 0.95, 5.3, 11.5, 0.9, "목표  →  논문 구조를 동작하는 시스템으로 재현 · 실데이터로 검증",
     18, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 7. 시스템 구성
s = prs.slides.add_slide(BLANK)
header(s, "03  구현 과정", "시스템 구성", 7)
blocks = [("데이터", ["12개 자산군 수익률", "거시지표", "ETF 실수익률 (추가)"]),
          ("백엔드", ["Python · FastAPI", "6단계 파이프라인", "NumPy · SciPy 계산"]),
          ("프론트엔드", ["React · TypeScript", "대시보드 4개 탭", "IPS 설정 · 재실행"]),
          ("배포 · 협업", ["GitHub 저장소", "Render 클라우드", "Pull Request 리뷰"])]
for i, (t, items) in enumerate(blocks):
    x = 0.6 + i * 3.14
    card(s, x, 1.9, 2.7, 3.0, t, items, accent=BLUE if i != 1 else NAVY, title_size=20, item_size=14)
    if i < 3:
        text(s, x + 2.72, 3.0, 0.4, 0.6, "→", 26, True, BLUE, PP_ALIGN.CENTER)
text(s, 0.6, 5.6, 12, 0.4, "계산 원칙", 14, True, NAVY)
text(s, 0.6, 6.0, 12, 0.8, ["·  수치 계산: 결정론적 코드 (재현 가능)",
                            "·  해설 · 토론: LLM 또는 시뮬레이션 모드 (API 키 없이 동작)"], 13, color=BODY)

# ---------------------------------------------------------------- 8. 진행 현황
s = prs.slides.add_slide(BLANK)
header(s, "03  구현 현황", "진행 상태", 8)
rows = [("1~2단계", "거시 진단 · CMA · OAS 공분산", "완료", GREEN),
        ("3단계", "포트폴리오 후보 최적화 (MV · BL 등)", "완료", GREEN),
        ("4단계", "CRO 몬테카를로 CVaR · 거부 루프", "완료", GREEN),
        ("5~6단계", "위원회 투표 · CIO 메모", "완료", GREEN),
        ("대시보드", "4개 분석 탭 · IPS 설정 · 한국어화", "완료", GREEN),
        ("배포", "Render 상시 접속 주소", "PR 검토 중", AMBER),
        ("백테스트", "18개 ETF · 159개월 실데이터", "PR 검토 중", AMBER),
        ("LLM 연동", "실제 API 기반 에이전트 추론", "예정", GREY)]
text(s, 0.6, 1.65, 2, 0.3, "구분", 12, True, MUTED)
text(s, 2.8, 1.65, 6, 0.3, "내용", 12, True, MUTED)
text(s, 10.6, 1.65, 2, 0.3, "상태", 12, True, MUTED)
for i, (a, b, st, col) in enumerate(rows):
    y = 2.05 + i * 0.6
    if i % 2 == 0:
        box(s, 0.6, y - 0.08, 12.13, 0.56, fill=GRAY_BG)
    text(s, 0.8, y - 0.08, 2, 0.56, a, 14, True, NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=0)
    text(s, 2.8, y - 0.08, 7.5, 0.56, b, 14, color=BODY, anchor=MSO_ANCHOR.MIDDLE, spacing=0)
    box(s, 10.6, y - 0.01, 1.7, 0.4, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 10.6, y - 0.01, 1.7, 0.4, st, 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 9. 중간 결과
s = prs.slides.add_slide(BLANK)
header(s, "03  중간 결과", "기본 IPS 기준 파이프라인 산출물", 9)
kpis = [("기대수익률", "6.66%"), ("변동성", "9.65%"), ("샤프비율", "0.30"), ("95% CVaR", "13.78%")]
for i, (k, v) in enumerate(kpis):
    x = 0.6 + i * 3.08
    box(s, x, 1.8, 2.88, 1.9, fill=BLUE_BG if i == 0 else GRAY_BG)
    text(s, x, 2.05, 2.88, 0.4, k, 14, True, BODY, PP_ALIGN.CENTER)
    text(s, x, 2.55, 2.88, 0.9, v, 38, True, BLUE if i == 0 else NAVY, PP_ALIGN.CENTER)
items9 = [("채택안", "Black-Litterman 후보"),
          ("위원회", "찬성 11 · 반대 1 (크레딧 과열 우려)"),
          ("CIO 조정", "미국 단기국채 +2%p"),
          ("IPS 시연", "변동성 한도 축소 → CRO 거부 → 재최적화")]
for i, (k, v) in enumerate(items9):
    y = 4.15 + i * 0.55
    text(s, 0.8, y, 2.2, 0.4, k, 14, True, NAVY)
    text(s, 3.0, y, 9.5, 0.4, v, 14, color=BODY)
text(s, 0.6, 6.55, 12, 0.4, "※ 시뮬레이션 데이터 기준 · 백테스트 PR 반영 시 수치 변경 예정", 11, color=MUTED)

# ---------------------------------------------------------------- 10. 이슈
s = prs.slides.add_slide(BLANK)
header(s, "04  향후 계획", "현재 이슈 및 과제", 10)
issues = [("IPS 기준 재설정", ["기본 변동성 밴드 8~12%", "21개 후보 중 적격 2개", "→ 기준값 재검토"]),
          ("데이터 신뢰성", ["현재: 시뮬레이션 수익률", "실데이터 ETF 전환 필요", "→ 결과 수치 재산출"]),
          ("LLM 실연동", ["현재: 시뮬레이션 모드", "에이전트 추론 품질 검증 필요", "→ API 연동 · 비용 관리"]),
          ("성과 검증", ["백테스트 기초 구현", "거래비용 · 비교군 부족", "→ 강건성 분석 보강"])]
for i, (t, items) in enumerate(issues):
    x = 0.6 + i * 3.08
    card(s, x, 1.8, 2.88, 4.3, t, items, accent=AMBER, title_size=17, item_size=13)

# ---------------------------------------------------------------- 11. 로드맵
s = prs.slides.add_slide(BLANK)
header(s, "04  향후 계획", "기말까지 로드맵", 11)
phases = [("Phase 1", "1~2주", "통합 · 정비", ["팀원 PR 병합", "배포 주소 단일화", "IPS 기본값 조정"]),
          ("Phase 2", "3~4주", "실데이터 전환", ["ETF 수익률 기반 CMA", "공분산 재추정", "결과 재산출"]),
          ("Phase 3", "5~6주", "LLM 실연동", ["거시 · CIO 에이전트 우선", "프롬프트 · 출력 검증", "시뮬레이션 대비 비교"]),
          ("Phase 4", "7~8주", "검증 · 마무리", ["거래비용 · 민감도 분석", "60/40 · 동일가중 비교", "보고서 · 기말 발표"])]
box(s, 0.9, 2.35, 11.6, 0.04, fill=LINE)
for i, (ph, wk, t, items) in enumerate(phases):
    x = 0.6 + i * 3.08
    text(s, x + 0.25, 1.65, 2.6, 0.3, ph, 12, True, BLUE)
    text(s, x + 1.2, 1.6, 1.6, 0.4, wk, 16, True, NAVY)
    box(s, x + 0.25, 2.17, 0.4, 0.4, fill=BLUE if i == 0 else NAVY, shape=MSO_SHAPE.OVAL)
    box(s, x, 2.9, 2.88, 3.0, fill=BLUE_BG if i == 0 else GRAY_BG)
    text(s, x + 0.25, 3.15, 2.4, 0.5, t, 20, True, NAVY)
    text(s, x + 0.25, 3.9, 2.45, 2.0, ["·  " + it for it in items], 14, color=BODY, spacing=12)
text(s, 0.6, 6.55, 12, 0.4, "※ 주차는 중간발표 이후 기준", 11, color=MUTED)

# ---------------------------------------------------------------- 12. 기말 산출물 + Q&A
s = prs.slides.add_slide(BLANK)
header(s, "04  향후 계획", "기말 목표 산출물", 12)
outs = [("시스템", "실데이터 · LLM 연동 완성본"),
        ("라이브 데모", "상시 접속 웹 주소"),
        ("검증 결과", "백테스트 · 민감도 분석"),
        ("문서", "최종 보고서 · 코드 리뷰 가이드")]
for i, (k, v) in enumerate(outs):
    x = 0.6 + (i % 2) * 6.15
    y = 1.8 + (i // 2) * 1.75
    box(s, x, y, 5.95, 1.5, fill=GRAY_BG)
    box(s, x, y, 0.08, 1.5, fill=BLUE)
    text(s, x + 0.4, y + 0.3, 5.3, 0.5, k, 20, True, NAVY)
    text(s, x + 0.4, y + 0.85, 5.3, 0.4, v, 14, color=BODY)
box(s, 0.6, 5.6, 12.13, 1.1, fill=NAVY)
text(s, 0.6, 5.6, 12.13, 1.1, "Q & A", 30, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT)

